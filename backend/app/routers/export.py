from io import BytesIO

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from sqlalchemy import extract, select
from sqlalchemy.ext.asyncio import AsyncSession

from ..db import get_db
from ..models import LearningEvent, Participant, Satisfaction, Session


router = APIRouter(prefix="/export", tags=["export"])


@router.get("/excel")
async def export_excel(
    year: int | None = Query(default=None, ge=2020, le=2100),
    month: int | None = Query(default=None, ge=1, le=12),
    scenario: str | None = None,
    db: AsyncSession = Depends(get_db),
):
    stmt = select(Session).order_by(Session.started_at)
    if year:
        stmt = stmt.where(extract("year", Session.started_at) == year)
    if month:
        stmt = stmt.where(extract("month", Session.started_at) == month)
    if scenario:
        stmt = stmt.where(Session.scenario_code == scenario)

    sessions = (await db.scalars(stmt)).all()
    ids = [s.id for s in sessions]

    participants = []
    events = []
    sats = []
    if ids:
        participants = (await db.scalars(select(Participant).where(Participant.session_id.in_(ids)))).all()
        events = (await db.scalars(select(LearningEvent).where(LearningEvent.session_id.in_(ids)).order_by(LearningEvent.created_at))).all()
        sats = (await db.scalars(select(Satisfaction).where(Satisfaction.session_id.in_(ids)))).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Sessions"
    ws.append(["session_id","session_code","scenario","version","participant_count","started_at","completed_at","duration_seconds","total_score","completed"])
    for s in sessions:
        ws.append([str(s.id),s.session_code,s.scenario_code,s.scenario_version,s.participant_count,s.started_at,s.completed_at,s.duration_seconds,s.total_score,s.completed])

    ws2 = wb.create_sheet("Participants")
    ws2.append(["participant_id","session_id","member_no","profession","role"])
    for p in participants:
        ws2.append([str(p.id),str(p.session_id),p.member_no,p.profession,p.role])

    ws3 = wb.create_sheet("Learning_Events")
    ws3.append(["event_id","session_id","event_type","step","question_id","selected_option","is_correct","attempt_number","elapsed_ms","created_at"])
    for e in events:
        ws3.append([str(e.id),str(e.session_id),e.event_type,e.step,e.question_id,e.selected_option,e.is_correct,e.attempt_number,e.elapsed_ms,e.created_at])

    ws4 = wb.create_sheet("Satisfaction")
    ws4.append(["session_id","q1","q2","q3","q4","q5","average","submitted_at"])
    for s in sats:
        ws4.append([str(s.session_id),s.q1,s.q2,s.q3,s.q4,s.q5,s.average,s.submitted_at])

    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)

    filename = f"Gener8_{scenario or 'all'}_{year or 'all'}_{month or 'all'}.xlsx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
